"""Deterministic bounded attachment extraction and local OCR derivative contract (MD-A3).

Extracts bounded text from verified MD-A2 quarantine attachments under `data/mail-desk/attachments/<run-id>/`.
Enforces strict character, page, slide, grid, and time limits.
For image-only or mixed PDFs, generates an isolated local OCR derivative under
`derivatives/<filename>.ocr.pdf` while strictly verifying workspace lock ownership,
sibling temp files, atomic no-clobber promotion, and pre/post source immutability.
"""

from __future__ import annotations

import concurrent.futures
import hashlib
import multiprocessing
import os
from pathlib import Path, PurePosixPath
import re
import signal
import subprocess
import sys
import time
from typing import Any, Callable
import uuid

def _init_win32_signatures(kernel32: Any) -> None:
    """Explicitly declare 64-bit safe argtypes and restypes for Win32 Job and Process APIs.

    Ensures that HANDLE returns and parameters use pointer-wide types (wintypes.HANDLE / c_void_p),
    preventing 32-bit integer truncation or pointer sign-extension on 64-bit Windows.
    """
    if kernel32 is None:
        return
    import ctypes
    from ctypes import wintypes

    # CreateJobObjectW(lpJobAttributes, lpName) -> HANDLE
    kernel32.CreateJobObjectW.argtypes = [wintypes.LPVOID, wintypes.LPCWSTR]
    kernel32.CreateJobObjectW.restype = wintypes.HANDLE

    # SetInformationJobObject(hJob, JobObjectInformationClass, lpJobObjectInformation, cbJobObjectInformationLength) -> BOOL
    kernel32.SetInformationJobObject.argtypes = [
        wintypes.HANDLE,
        ctypes.c_int,
        wintypes.LPVOID,
        wintypes.DWORD,
    ]
    kernel32.SetInformationJobObject.restype = wintypes.BOOL

    # AssignProcessToJobObject(hJob, hProcess) -> BOOL
    kernel32.AssignProcessToJobObject.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
    kernel32.AssignProcessToJobObject.restype = wintypes.BOOL

    # TerminateJobObject(hJob, uExitCode) -> BOOL
    kernel32.TerminateJobObject.argtypes = [wintypes.HANDLE, wintypes.UINT]
    kernel32.TerminateJobObject.restype = wintypes.BOOL

    # OpenProcess(dwDesiredAccess, bInheritHandle, dwProcessId) -> HANDLE
    kernel32.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
    kernel32.OpenProcess.restype = wintypes.HANDLE

    # CloseHandle(hObject) -> BOOL
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL

    # GetExitCodeProcess(hProcess, lpExitCode) -> BOOL
    kernel32.GetExitCodeProcess.argtypes = [wintypes.HANDLE, wintypes.LPDWORD]
    kernel32.GetExitCodeProcess.restype = wintypes.BOOL


_kernel32: Any = None
_win32_init_error: str | None = None


def _get_verified_kernel32() -> Any:
    """Return the verified kernel32 handle with initialized ctypes signatures.

    Fails closed with RuntimeError if running on Windows and kernel32 or signature
    initialization is unavailable or failed. Never falls back to unconfigured windll.kernel32.
    """
    if sys.platform != "win32":
        raise RuntimeError("Win32 kernel32 API is not available on non-Windows platforms.")
    if _kernel32 is not None:
        return _kernel32
    err = _win32_init_error or "kernel32 signatures not initialized"
    raise RuntimeError(f"Win32 kernel32 confinement unavailable: {err}")


if sys.platform == "win32":
    try:
        import ctypes
        _k32 = ctypes.windll.kernel32
        _init_win32_signatures(_k32)
        _kernel32 = _k32
        _win32_init_error = None
    except Exception as exc:
        _kernel32 = None
        _win32_init_error = f"kernel32 signature initialization failed: {exc}"
else:
    _kernel32 = None
    _win32_init_error = "Non-Windows platform"

from core.common import resolve_data_dir
from core.attachment_policy import DEFAULT_ATTACHMENT_POLICY
from core.attachments import HashDriftError
from core.attachment_fetch import (
    ActiveContentBlockedError,
    DisallowedExtensionError,
    ExtensionMimeDriftError,
    MimeDriftError,
    QuarantineCollisionError,
    SymlinkEscapeError,
    WorkspaceLockError,
    _atomic_no_clobber_promote,
    check_quarantine_path_security,
    detect_mime_and_active_content,
    is_valid_run_id,
    validate_attachment_filename,
    validate_mime_and_extension,
    verify_workspace_lock,
)


# Default Fallback Limits
DEFAULT_MAX_CHARS_PER_ATTACHMENT = 15000
DEFAULT_MAX_PDF_PAGES = 10
DEFAULT_MAX_OCR_PAGES = 3
DEFAULT_MAX_DOCX_PARAGRAPHS = 40
DEFAULT_MAX_PPTX_SLIDES = 15
DEFAULT_MAX_XLSX_SHEETS = 2
DEFAULT_MAX_XLSX_ROWS = 50
DEFAULT_MAX_XLSX_COLS = 10
DEFAULT_PROCESS_TIMEOUT_SECONDS = 20
DEFAULT_OCR_TIMEOUT_SECONDS = 30


def _resolve_extraction_limits(policy: dict[str, Any] | None) -> dict[str, Any]:
    """Dynamically resolve extraction and budget limits from nested or flat policy with defaults."""
    pol = policy if isinstance(policy, dict) else DEFAULT_ATTACHMENT_POLICY
    ext_cfg = pol.get("extraction", {}) if isinstance(pol.get("extraction"), dict) else {}
    bud_cfg = pol.get("budget", {}) if isinstance(pol.get("budget"), dict) else {}
    trans_cfg = pol.get("transport", {}) if isinstance(pol.get("transport"), dict) else {}

    max_chars = bud_cfg.get(
        "max_chars_per_attachment",
        pol.get("max_chars_per_attachment", DEFAULT_MAX_CHARS_PER_ATTACHMENT),
    )
    max_pdf_pages = ext_cfg.get("pdf_max_pages", pol.get("max_pdf_pages", DEFAULT_MAX_PDF_PAGES))
    max_ocr_pages = ext_cfg.get("ocr_max_pages", pol.get("max_ocr_pages", DEFAULT_MAX_OCR_PAGES))
    process_timeout = ext_cfg.get(
        "process_timeout_seconds",
        ext_cfg.get(
            "extraction_timeout_seconds",
            pol.get("process_timeout_seconds", pol.get("extraction_timeout_seconds", DEFAULT_PROCESS_TIMEOUT_SECONDS)),
        ),
    )
    ocr_timeout = ext_cfg.get("ocr_timeout_seconds", pol.get("ocr_timeout_seconds", DEFAULT_OCR_TIMEOUT_SECONDS))
    max_docx_paragraphs = ext_cfg.get("docx_max_paragraphs", pol.get("max_docx_paragraphs", DEFAULT_MAX_DOCX_PARAGRAPHS))
    max_pptx_slides = ext_cfg.get("pptx_max_slides", pol.get("max_pptx_slides", DEFAULT_MAX_PPTX_SLIDES))
    max_xlsx_sheets = ext_cfg.get("xlsx_max_sheets", pol.get("max_xlsx_sheets", DEFAULT_MAX_XLSX_SHEETS))
    max_xlsx_rows = ext_cfg.get("xlsx_max_rows_per_sheet", pol.get("max_xlsx_rows", DEFAULT_MAX_XLSX_ROWS))
    max_xlsx_cols = ext_cfg.get("xlsx_max_cols_per_sheet", pol.get("max_xlsx_cols", DEFAULT_MAX_XLSX_COLS))

    disallowed_exts = trans_cfg.get(
        "disallowed_extensions",
        pol.get("disallowed_extensions", DEFAULT_ATTACHMENT_POLICY["transport"]["disallowed_extensions"]),
    )

    return {
        "max_chars_per_attachment": int(max_chars),
        "max_pdf_pages": int(max_pdf_pages),
        "max_ocr_pages": int(max_ocr_pages),
        "process_timeout_seconds": float(process_timeout),
        "ocr_timeout_seconds": float(ocr_timeout),
        "max_docx_paragraphs": int(max_docx_paragraphs),
        "max_pptx_slides": int(max_pptx_slides),
        "max_xlsx_sheets": int(max_xlsx_sheets),
        "max_xlsx_rows": int(max_xlsx_rows),
        "max_xlsx_cols": int(max_xlsx_cols),
        "disallowed_extensions": set(str(e).lower() for e in disallowed_exts),
    }


def is_process_alive(pid: int) -> bool:
    """Check if a process with the given PID is actively running across OS platforms."""
    if pid <= 0:
        return False
    if sys.platform == "win32":
        kernel32 = _get_verified_kernel32()
        import ctypes
        from ctypes import wintypes
        PROCESS_QUERY_LIMITED_INFORMATION = 0x1000
        STILL_ACTIVE = 259
        h = kernel32.OpenProcess(PROCESS_QUERY_LIMITED_INFORMATION, False, pid)
        if not h:
            return False
        exit_code = wintypes.DWORD()
        try:
            if kernel32.GetExitCodeProcess(h, ctypes.byref(exit_code)):
                return exit_code.value == STILL_ACTIVE
            return False
        finally:
            try:
                kernel32.CloseHandle(h)
            except Exception:
                pass
    else:
        try:
            os.kill(pid, 0)
            return True
        except (OSError, ProcessLookupError):
            return False


class WindowsJobObject:
    """Encapsulates a Windows Job Object for automatic process tree confinement."""

    def __init__(self) -> None:
        self.handle = None
        if sys.platform == "win32":
            kernel32 = _get_verified_kernel32()
            import ctypes
            h = kernel32.CreateJobObjectW(None, None)
            if not h:
                err = ctypes.GetLastError()
                raise RuntimeError(f"CreateJobObjectW failed with error code {err}")
            self.handle = h

            # Configure JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE (0x2000)
            class JOBOBJECT_BASIC_LIMIT_INFORMATION(ctypes.Structure):
                _fields_ = [
                    ("PerProcessUserTimeLimit", ctypes.c_int64),
                    ("PerJobUserTimeLimit", ctypes.c_int64),
                    ("LimitFlags", ctypes.c_uint32),
                    ("MinimumWorkingSetSize", ctypes.c_size_t),
                    ("MaximumWorkingSetSize", ctypes.c_size_t),
                    ("ActiveProcessLimit", ctypes.c_uint32),
                    ("Affinity", ctypes.c_size_t),
                    ("PriorityClass", ctypes.c_uint32),
                    ("SchedulingClass", ctypes.c_uint32),
                ]

            class IO_COUNTERS(ctypes.Structure):
                _fields_ = [
                    ("ReadOperationCount", ctypes.c_uint64),
                    ("WriteOperationCount", ctypes.c_uint64),
                    ("OtherOperationCount", ctypes.c_uint64),
                    ("ReadTransferCount", ctypes.c_uint64),
                    ("WriteTransferCount", ctypes.c_uint64),
                    ("OtherTransferCount", ctypes.c_uint64),
                ]

            class JOBOBJECT_EXTENDED_LIMIT_INFORMATION(ctypes.Structure):
                _fields_ = [
                    ("BasicLimitInformation", JOBOBJECT_BASIC_LIMIT_INFORMATION),
                    ("IoInfo", IO_COUNTERS),
                    ("ProcessMemoryLimit", ctypes.c_size_t),
                    ("JobMemoryLimit", ctypes.c_size_t),
                    ("PeakProcessMemoryLimit", ctypes.c_size_t),
                    ("PeakJobMemoryLimit", ctypes.c_size_t),
                ]

            info = JOBOBJECT_EXTENDED_LIMIT_INFORMATION()
            info.BasicLimitInformation.LimitFlags = 0x2000  # JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
            JobObjectExtendedLimitInformation = 9
            res = kernel32.SetInformationJobObject(
                self.handle,
                JobObjectExtendedLimitInformation,
                ctypes.byref(info),
                ctypes.sizeof(info),
            )
            if not res:
                err = ctypes.GetLastError()
                try:
                    kernel32.CloseHandle(self.handle)
                except Exception:
                    pass
                self.handle = None
                raise RuntimeError(f"SetInformationJobObject failed with error code {err}")

    def assign(self, pid: int) -> None:
        if sys.platform != "win32":
            return
        if not self.handle:
            raise RuntimeError("Cannot assign process: Job Object handle is invalid")
        kernel32 = _get_verified_kernel32()
        import ctypes
        PROCESS_ALL_ACCESS = 0x1F0FFF
        h_proc = kernel32.OpenProcess(PROCESS_ALL_ACCESS, False, pid)
        if not h_proc:
            err = ctypes.GetLastError()
            raise RuntimeError(f"OpenProcess for PID {pid} failed with error code {err}")
        try:
            res = kernel32.AssignProcessToJobObject(self.handle, h_proc)
            if not res:
                err = ctypes.GetLastError()
                raise RuntimeError(f"AssignProcessToJobObject failed for PID {pid} with error code {err}")
        finally:
            try:
                kernel32.CloseHandle(h_proc)
            except Exception:
                pass

    def terminate(self) -> None:
        if self.handle and sys.platform == "win32":
            try:
                kernel32 = _get_verified_kernel32()
                kernel32.TerminateJobObject(self.handle, 1)
            except Exception:
                pass
            try:
                kernel32 = _get_verified_kernel32()
                kernel32.CloseHandle(self.handle)
            except Exception:
                pass
            self.handle = None

    def close(self) -> None:
        if self.handle and sys.platform == "win32":
            try:
                kernel32 = _get_verified_kernel32()
                kernel32.CloseHandle(self.handle)
            except Exception:
                pass
            self.handle = None


def _terminate_single_process(proc: Any) -> None:
    """Safely terminate only the single process without touching any process group."""
    try:
        proc.kill()
    except Exception:
        pass
    try:
        proc.join(timeout=2.0)
    except Exception:
        pass


def terminate_process_tree(proc: Any, job: Any = None) -> None:
    """Forcefully terminate an entire process tree including all child/grandchild processes.

    On Windows:
      1. Terminates via Win32 Job Object (if configured), terminating all descendants kernel-side.
      2. Executes 'taskkill /F /T /PID <pid>' as secondary defense for detached child process trees.
      3. Issues proc.kill() on the Python Process object and joins with timeout.
    On POSIX:
      1. Verifies that the target process is indeed the leader of its own process group (os.getpgid(pid) == pid).
         Only if confirmed does it kill the process group via os.killpg(pid, signal.SIGKILL),
         guaranteeing that a shared or parent process group is NEVER killed!
      2. Issues proc.kill() on the Python Process object and joins with timeout.
    """
    pid = getattr(proc, "pid", None)
    if not pid:
        return

    # 1. Windows Job Object termination (guaranteed kernel-side tree kill)
    if sys.platform == "win32" and job is not None:
        try:
            job.terminate()
        except Exception:
            pass

    # 2. Windows taskkill tree-kill (secondary defense)
    if sys.platform == "win32":
        try:
            subprocess.run(
                ["taskkill", "/F", "/T", "/PID", str(pid)],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=5.0,
                check=False,
            )
        except Exception:
            pass
    else:
        # POSIX process group termination: strictly only if process is its own group leader
        try:
            if hasattr(os, "getpgid") and hasattr(os, "killpg"):
                pgid = os.getpgid(pid)
                if pgid == pid:
                    sig_kill = getattr(signal, "SIGKILL", 9)
                    os.killpg(pgid, sig_kill)
        except Exception:
            pass

    # 3. Direct process kill and join
    _terminate_single_process(proc)


def _process_worker_target(child_conn: Any, func: Callable[..., Any], args: tuple, kwargs: dict) -> None:
    """Entry point executed inside an isolated worker process with parent authorization handshake."""
    # Under POSIX, isolate process into its own process group (pgid == pid).
    # If setpgid fails, worker MUST NOT report READY and MUST NOT execute payload!
    if sys.platform != "win32":
        try:
            if hasattr(os, "setpgid"):
                os.setpgid(0, 0)
            else:
                raise OSError("os.setpgid is not available on this platform")
        except Exception as exc:
            try:
                child_conn.send(f"ERROR:setpgid_failed:{exc}")
            except Exception:
                pass
            try:
                child_conn.close()
            except Exception:
                pass
            return

    # Step 1 of Handshake: Signal READY to parent
    try:
        child_conn.send("READY")
    except Exception:
        try:
            child_conn.close()
        except Exception:
            pass
        return

    # Step 2 of Handshake: Worker MUST block and wait for parent authorization before executing func.
    # The worker MUST NOT run the target function before the parent establishes confinement (e.g. Job Object or verified PGID).
    try:
        if child_conn.poll(10.0):
            msg = child_conn.recv()
            if msg != "START":
                try:
                    child_conn.close()
                except Exception:
                    pass
                return
        else:
            try:
                child_conn.close()
            except Exception:
                pass
            return
    except Exception:
        try:
            child_conn.close()
        except Exception:
            pass
        return

    # Step 3 of Handshake: Confinement verified; execute payload
    try:
        res = func(*args, **kwargs)
        child_conn.send((True, res))
    except Exception as exc:
        child_conn.send((False, exc))
    finally:
        try:
            child_conn.close()
        except Exception:
            pass


def _run_with_thread_fallback(
    func: Callable[..., Any],
    args: tuple,
    kwargs: dict,
    timeout_seconds: float,
) -> Any:
    """Explicit test-internal helper fallback for unpicklable callables in unit tests."""
    executor = concurrent.futures.ThreadPoolExecutor(max_workers=1)
    fut = executor.submit(func, *args, **kwargs)
    try:
        res = fut.result(timeout=timeout_seconds)
        executor.shutdown(wait=False)
        return res
    except concurrent.futures.TimeoutError as err:
        executor.shutdown(wait=False, cancel_futures=True)
        raise TimeoutError(f"Execution timed out after {timeout_seconds} seconds") from err
    except Exception:
        executor.shutdown(wait=False)
        raise


def run_with_timeout(
    func: Callable[..., Any],
    args: tuple = (),
    kwargs: dict | None = None,
    timeout_seconds: float = 20.0,
    allow_thread_fallback: bool = False,
) -> Any:
    """Execute a callable with an interruptible timeout and process tree confinement.

    Executes inside an isolated multiprocessing.Process worker attached to an OS process
    group (POSIX) or Win32 Job Object (Windows).
    Under Windows, a two-phase handshake ensures the worker does NOT begin execution until
    the Job Object is successfully created, configured, and assigned.
    On timeout, the entire process tree (including all descendants such as Tesseract or
    Ghostscript) is forcefully terminated via tree-kill.
    The production path strictly enforces process isolation and fails closed if
    the callable or arguments are not picklable.
    """
    if kwargs is None:
        kwargs = {}

    can_pickle = False
    try:
        import pickle
        pickle.dumps((func, args, kwargs))
        can_pickle = True
    except Exception as e:
        if not allow_thread_fallback:
            raise RuntimeError(
                f"Process isolation failed: callable or arguments are not picklable ({e}). "
                f"Thread fallback is forbidden in production path."
            ) from e
        return _run_with_thread_fallback(func, args, kwargs, timeout_seconds)

    job: WindowsJobObject | None = None
    parent_conn, child_conn = multiprocessing.Pipe(duplex=True)
    proc = multiprocessing.Process(
        target=_process_worker_target,
        args=(child_conn, func, args, kwargs),
    )

    try:
        proc.start()
        child_conn.close()

        # Step 1: Wait for READY signal from worker
        if not parent_conn.poll(5.0):
            _terminate_single_process(proc)
            parent_conn.close()
            raise RuntimeError("Worker process failed to send READY handshake signal.")

        try:
            ready_msg = parent_conn.recv()
        except (EOFError, BrokenPipeError) as e:
            _terminate_single_process(proc)
            parent_conn.close()
            raise RuntimeError(f"Worker process exited before sending READY handshake signal: {e}") from e

        if ready_msg != "READY":
            _terminate_single_process(proc)
            parent_conn.close()
            raise RuntimeError(f"Unexpected handshake message from worker: {ready_msg}")

        # Step 2: Platform-specific confinement verification BEFORE releasing worker
        if sys.platform == "win32":
            try:
                job = WindowsJobObject()
                job.assign(proc.pid)
            except Exception as e:
                # Confinement setup failed! Terminate worker before giving release; fail closed!
                try:
                    parent_conn.send("ABORT")
                except Exception:
                    pass
                terminate_process_tree(proc, job=job)
                if job:
                    job.close()
                try:
                    parent_conn.close()
                except Exception:
                    pass
                raise RuntimeError(
                    f"Windows Job Object confinement failed: {e}. Worker terminated before execution."
                ) from e
        else:
            # Under POSIX: verify worker is confirmed leader of its own process group (os.getpgid(proc.pid) == proc.pid)
            try:
                if not hasattr(os, "getpgid"):
                    raise OSError("os.getpgid is not available on this platform")
                actual_pgid = os.getpgid(proc.pid)
                if actual_pgid != proc.pid:
                    raise RuntimeError(
                        f"Process group mismatch: expected PGID {proc.pid}, got {actual_pgid}"
                    )
            except Exception as e:
                # Confinement verification failed! Terminate worker strictly individually; NEVER kill unconfirmed process group!
                try:
                    parent_conn.send("ABORT")
                except Exception:
                    pass
                _terminate_single_process(proc)
                try:
                    parent_conn.close()
                except Exception:
                    pass
                raise RuntimeError(
                    f"POSIX process group confinement failed: {e}. Worker terminated before execution."
                ) from e

        # Step 3: Confinement verified! Authorize worker to begin execution
        parent_conn.send("START")

    except Exception:
        if proc.is_alive():
            if sys.platform == "win32" and job is not None:
                terminate_process_tree(proc, job=job)
            else:
                _terminate_single_process(proc)
        if job:
            job.close()
        try:
            parent_conn.close()
        except Exception:
            pass
        raise

    has_data = False
    try:
        has_data = parent_conn.poll(timeout_seconds)
    except Exception:
        has_data = False

    if not has_data:
        # Terminate entire process tree (worker and all descendants)
        terminate_process_tree(proc, job=job)
        if job:
            job.close()
        try:
            parent_conn.close()
        except Exception:
            pass
        raise TimeoutError(
            f"Execution timed out after {timeout_seconds} seconds (process tree for PID {proc.pid} terminated)."
        )

    try:
        ok, res = parent_conn.recv()
    except Exception as e:
        terminate_process_tree(proc, job=job)
        if job:
            job.close()
        parent_conn.close()
        raise RuntimeError(f"Worker process communication failed: {e}") from e

    proc.join(timeout=2.0)
    if job:
        job.close()
    parent_conn.close()
    if ok:
        return res
    raise res


def _cleanup_temp_artifacts(known_temp_path: Path | None, *args: Any, **kwargs: Any) -> None:
    """Safely unlink exclusively the parent-owned, invocation-specific temp derivative path.

    Never globs or removes sibling temp files belonging to concurrent invocations under the same stem.
    Any extra arguments are ignored for backward-compatibility.
    """
    if known_temp_path is not None:
        try:
            if known_temp_path.is_symlink() or known_temp_path.exists():
                known_temp_path.unlink(missing_ok=True)
        except Exception:
            pass


def _get_tool_version(tool_name: str) -> str:
    """Retrieve genuine tool/package version dynamically."""
    try:
        if tool_name == "pymupdf":
            import pymupdf
            return str(getattr(pymupdf, "__version__", getattr(pymupdf, "version", "unknown")))
        elif tool_name == "python-docx":
            import docx
            return str(getattr(docx, "__version__", "unknown"))
        elif tool_name == "openpyxl":
            import openpyxl
            return str(getattr(openpyxl, "__version__", "unknown"))
        elif tool_name == "python-pptx":
            import pptx
            return str(getattr(pptx, "__version__", "unknown"))
        elif tool_name == "ocrmypdf":
            import ocrmypdf
            return str(getattr(ocrmypdf, "__version__", "unknown"))
        elif tool_name == "markitdown":
            import markitdown
            return str(getattr(markitdown, "__version__", "unknown"))
        elif tool_name == "built_in":
            return f"python-{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}"
    except Exception:
        pass
    return "unknown"


def validate_mda2_fetch_result(
    fetch_result: dict[str, Any],
    expected_sha256: str,
) -> tuple[str, str, str, str]:
    """Validate that fetch_result strictly conforms to the verified MD-A2 envelope contract.

    Parameters:
        fetch_result: Dictionary returned from MD-A2 fetch operation.
        expected_sha256: Mandatory 64-character SHA-256 hash expected for this attachment.

    Returns:
        (run_id, relative_path, verified_sha256, effective_mime_type)
    Raises:
        ValueError: On missing/invalid envelope keys, status, run_id, or empty expected_sha256.
        HashDriftError: On hash drift between fetch_sha256, inventory_sha256, or expected_sha256.
    """
    if not isinstance(fetch_result, dict):
        raise ValueError("fetch_result must be a valid dictionary.")

    if not expected_sha256 or not isinstance(expected_sha256, str) or not expected_sha256.strip():
        raise ValueError("expected_sha256 is mandatory and cannot be empty.")

    exp_norm = expected_sha256.strip().lower()
    sha_pattern = re.compile(r"^[0-9a-f]{64}$")
    if not sha_pattern.fullmatch(exp_norm):
        raise ValueError(f"expected_sha256 is not a valid 64-char hex SHA-256: {expected_sha256!r}")

    status = fetch_result.get("status")
    if status not in ("fetched", "already_fetched"):
        raise ValueError(f"Invalid or unverified MD-A2 fetch status: {status!r}")

    run_id = fetch_result.get("run_id")
    if not run_id or not isinstance(run_id, str) or not is_valid_run_id(run_id):
        raise ValueError(f"Invalid or unsafe MD-A2 run_id: {run_id!r}")

    rel_path = fetch_result.get("relative_path")
    if not rel_path or not isinstance(rel_path, str) or not rel_path.strip():
        raise ValueError("MD-A2 fetch_result missing 'relative_path'.")

    eff_mime = fetch_result.get("effective_mime_type")
    if not eff_mime or not isinstance(eff_mime, str) or not eff_mime.strip():
        raise ValueError("MD-A2 fetch_result missing 'effective_mime_type'.")

    fetch_sha = fetch_result.get("fetch_sha256")
    inv_sha = fetch_result.get("inventory_sha256")
    if not fetch_sha or not inv_sha:
        raise ValueError("MD-A2 fetch_result must contain both 'fetch_sha256' and 'inventory_sha256'.")

    fetch_sha_norm = str(fetch_sha).strip().lower()
    inv_sha_norm = str(inv_sha).strip().lower()

    if not sha_pattern.fullmatch(fetch_sha_norm):
        raise ValueError(f"MD-A2 'fetch_sha256' is not a valid 64-char hex SHA-256: {fetch_sha!r}")
    if not sha_pattern.fullmatch(inv_sha_norm):
        raise ValueError(f"MD-A2 'inventory_sha256' is not a valid 64-char hex SHA-256: {inv_sha!r}")

    if fetch_sha_norm != inv_sha_norm:
        raise HashDriftError(
            f"MD-A2 hash drift: fetch_sha256 '{fetch_sha_norm}' != inventory_sha256 '{inv_sha_norm}'"
        )

    if exp_norm != fetch_sha_norm:
        raise HashDriftError(
            f"Expected SHA-256 '{exp_norm}' does not match MD-A2 fetch_sha256 '{fetch_sha_norm}'"
        )

    return run_id.strip(), rel_path.strip(), fetch_sha_norm, eff_mime.strip().lower()


def resolve_quarantine_source_file(
    run_id: str,
    relative_path: str,
    base_data_dir: Path,
) -> tuple[Path, Path, Path]:
    """Resolve and strictly validate containment of quarantine file under data/mail-desk/attachments/<run-id>/.

    Preserves unresolved paths for symlink/junction inspection before resolving for containment.

    Returns:
        (target_file, run_dir, raw_attachments_root)
    Raises:
        ValueError: On absolute path, traversal sequence, or escaped run-id boundary.
        SymlinkEscapeError: On symlink or Windows reparse point traversal.
    """
    clean_path = str(relative_path).replace("\\", "/")

    # Block absolute paths and Windows drive letters
    if clean_path.startswith("/") or ":" in clean_path:
        raise ValueError(f"Absolute or drive-qualified relative_path rejected: {relative_path!r}")

    # Block directory traversal sequences
    parts = clean_path.split("/")
    if ".." in parts:
        raise ValueError(f"Path traversal sequence '..' rejected in relative_path: {relative_path!r}")

    if clean_path.startswith("data/mail-desk/"):
        clean_path = clean_path[len("data/mail-desk/"):]

    expected_prefix = f"attachments/{run_id}/"
    if not clean_path.startswith(expected_prefix):
        raise ValueError(
            f"relative_path '{relative_path}' does not point into expected active run quarantine 'attachments/{run_id}/'"
        )

    file_inside = clean_path[len(expected_prefix):]
    if "/" in file_inside:
        raise ValueError(f"Subdirectories inside run quarantine rejected: {file_inside!r}")

    clean_filename = validate_attachment_filename(file_inside)

    raw_attachments_root = base_data_dir / "attachments"
    raw_run_dir = raw_attachments_root / run_id
    raw_target_file = raw_run_dir / clean_filename

    # Check symlinks/junctions on UNRESOLVED paths first
    check_quarantine_path_security(raw_target_file, raw_attachments_root)

    # Now resolve for containment verification
    attachments_root = raw_attachments_root.resolve()
    run_dir = raw_run_dir.resolve()
    target_file = raw_target_file.resolve()

    try:
        run_dir.relative_to(attachments_root)
    except ValueError:
        raise ValueError(f"Run directory '{run_dir}' escapes attachments root '{attachments_root}'")

    try:
        target_file.relative_to(run_dir)
    except ValueError:
        raise ValueError(f"Target file '{target_file}' escapes run directory '{run_dir}'")

    return target_file, run_dir, raw_attachments_root


def _run_ocr_derivative(
    source_pdf_path: Path,
    derivative_pdf_path: Path,
    max_pages: int = DEFAULT_MAX_OCR_PAGES,
    pages_to_ocr: list[int] | str | None = None,
    timeout_seconds: float = DEFAULT_OCR_TIMEOUT_SECONDS,
) -> tuple[bytes, str, int]:
    """Perform bounded local OCR on an image PDF, writing exclusively to the derivative path."""
    import pymupdf

    derivative_pdf_path.parent.mkdir(parents=True, exist_ok=True)

    if pages_to_ocr is not None:
        if isinstance(pages_to_ocr, (list, tuple)):
            pages_arg = ",".join(str(p) for p in pages_to_ocr)
        else:
            pages_arg = str(pages_to_ocr)
    else:
        pages_arg = f"1-{max_pages}"

    try:
        import ocrmypdf

        ocrmypdf.ocr(
            source_pdf_path,
            derivative_pdf_path,
            pages=pages_arg,
            force_ocr=True,
            output_type="pdf",
            optimize=0,
            fast_web_view=0,
        )

        if not derivative_pdf_path.exists():
            raise RuntimeError("OCR tool completed without writing output PDF.")

        deriv_bytes = derivative_pdf_path.read_bytes()
        doc = pymupdf.open(derivative_pdf_path)
        texts = [p.get_text() for p in doc]
        doc.close()
        return deriv_bytes, "\n".join(texts), len(texts)
    except Exception as e:
        derivative_pdf_path.unlink(missing_ok=True)
        raise RuntimeError(f"OCR tool execution failed or unavailable: {e}") from e


def _extract_content_internal(
    target_file: Path,
    actual_sha: str,
    filename: str,
    ext: str,
    eff_mime: str,
    run_dir: Path,
    run_id: str,
    raw_attachments_root: Path,
    base_data_dir: Path,
    limits: dict[str, Any],
    workspace_root: str | Path | None = None,
    lease_id: str | None = None,
    conversation_id: str | None = None,
    temp_deriv_path: Path | None = None,
    ocr_runner: Callable[..., Any] | None = None,
    lock_verifier: Callable[..., Any] | None = None,
) -> dict[str, Any]:
    """Perform bounded extraction under isolation."""
    file_bytes = target_file.read_bytes()
    max_chars = limits["max_chars_per_attachment"]
    max_pdf_pages = limits["max_pdf_pages"]
    max_ocr_pages = limits["max_ocr_pages"]
    ocr_timeout = limits["ocr_timeout_seconds"]

    status = "extracted"
    method = "native"
    tool = "built_in"
    tool_version = _get_tool_version("built_in")
    quality = "high"
    truncation_reason: str | None = None
    source_character_count: int | None = None
    extracted_text = ""
    derivative_sha256: str | None = None
    derivative_relative_path: str | None = None
    scope: dict[str, Any] = {
        "pages_processed": None,
        "pages_total": None,
        "ocr_pages": 0,
        "paragraphs": None,
        "slides": None,
        "sheets": None,
    }

    # 1. Plain Text / CSV / Markdown / Log / JSON
    if ext in [".txt", ".csv", ".tsv", ".md", ".log", ".json"] or eff_mime.startswith("text/"):
        decoded_str = ""
        for enc in ["utf-8", "cp1252", "latin1"]:
            try:
                decoded_str = file_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        source_character_count = len(decoded_str)
        if len(decoded_str) > max_chars:
            extracted_text = decoded_str[:max_chars]
            truncation_reason = "max_chars_exceeded"
        else:
            extracted_text = decoded_str
        method = "native"
        tool = "built_in"
        tool_version = _get_tool_version("built_in")

    # 2. PDF Documents
    elif ext == ".pdf" or eff_mime == "application/pdf":
        import pymupdf
        try:
            doc = pymupdf.open(target_file)
        except Exception as e:
            return {
                "status": "corrupt_attachment",
                "source_sha256": actual_sha,
                "derivative_sha256": None,
                "derivative_relative_path": None,
                "method": "native",
                "tool": "pymupdf",
                "tool_version": _get_tool_version("pymupdf"),
                "quality": "low",
                "scope": scope,
                "truncation_reason": None,
                "character_count": 0,
                "text": "",
                "error": f"PDF parse error: {e}",
            }

        total_pages = len(doc)
        scope["pages_total"] = total_pages
        pages_to_process = min(total_pages, max_pdf_pages)
        scope["pages_processed"] = pages_to_process
        if total_pages > max_pdf_pages:
            truncation_reason = "max_pages_exceeded"

        pages_info: list[tuple[int, bool, str, bool]] = []
        for p_idx in range(pages_to_process):
            page = doc[p_idx]
            p_text = page.get_text()
            clean_txt = p_text.strip() if p_text else ""
            has_txt = bool(clean_txt)
            images = page.get_images()
            has_imgs = bool(images)
            pages_info.append((p_idx + 1, has_txt, p_text, has_imgs))
        doc.close()

        digital_pages = [p for p in pages_info if p[1]]
        image_pages = [p for p in pages_info if not p[1]]

        # Case 2a: Pure Digital PDF
        if len(image_pages) == 0:
            method = "native"
            tool = "pymupdf"
            tool_version = _get_tool_version("pymupdf")
            quality = "high"
            scope["ocr_pages"] = 0

            lines = []
            cur_chars = 0
            for _, _, p_txt, _ in digital_pages:
                if cur_chars + len(p_txt) > max_chars:
                    rem = max_chars - cur_chars
                    lines.append(p_txt[:rem])
                    cur_chars += rem
                    truncation_reason = "max_chars_exceeded"
                    break
                lines.append(p_txt)
                cur_chars += len(p_txt)
            extracted_text = "\n".join(lines)
            if total_pages <= max_pdf_pages:
                source_character_count = len("\n".join(p_txt for _, _, p_txt, _ in digital_pages))
            else:
                source_character_count = None

        # Case 2b: Pure Image / Scanned PDF
        elif len(digital_pages) == 0:
            # 1. Lock & Path check immediately before temp creation
            active_lock_verifier = lock_verifier or verify_workspace_lock
            active_lock_verifier(
                workspace_root=workspace_root,
                lease_id=lease_id,
                conversation_id=conversation_id,
                data_dir=base_data_dir,
            )

            raw_derivatives_dir = run_dir / "derivatives"
            check_quarantine_path_security(raw_derivatives_dir, raw_attachments_root)
            derivatives_dir = raw_derivatives_dir.resolve()
            try:
                derivatives_dir.relative_to(run_dir)
            except ValueError:
                raise ValueError(f"Derivatives dir '{derivatives_dir}' escapes run directory '{run_dir}'")
            derivatives_dir.mkdir(parents=True, exist_ok=True)

            deriv_filename = f"{target_file.stem}.ocr.pdf"
            raw_deriv_path = derivatives_dir / deriv_filename
            check_quarantine_path_security(raw_deriv_path, raw_attachments_root)
            deriv_path = raw_deriv_path.resolve()
            try:
                deriv_path.relative_to(derivatives_dir)
            except ValueError:
                raise ValueError(f"Derivative path '{deriv_path}' escapes derivatives dir '{derivatives_dir}'")

            num_ocr_pages = min(pages_to_process, max_ocr_pages)
            pages_arg = list(range(1, num_ocr_pages + 1))

            pre_source_sha = hashlib.sha256(target_file.read_bytes()).hexdigest().lower()
            if pre_source_sha != actual_sha:
                raise HashDriftError("Source file hash changed prior to OCR execution.")

            temp_deriv = (
                temp_deriv_path
                if temp_deriv_path is not None
                else derivatives_dir / f".{deriv_filename}.{uuid.uuid4().hex}.tmp"
            )
            check_quarantine_path_security(temp_deriv, raw_attachments_root)
            was_newly_created = False

            active_ocr_runner = ocr_runner or _run_ocr_derivative
            try:
                deriv_bytes, ocr_text, ocr_count = run_with_timeout(
                    active_ocr_runner,
                    args=(target_file, temp_deriv, max_ocr_pages, pages_arg, ocr_timeout),
                    timeout_seconds=ocr_timeout,
                    allow_thread_fallback=False,
                )
                if not temp_deriv.exists():
                    temp_deriv.write_bytes(deriv_bytes)
                deriv_sha = hashlib.sha256(temp_deriv.read_bytes()).hexdigest().lower()

                # 2. Lock & Path check immediately before promotion
                active_lock_verifier(
                    workspace_root=workspace_root,
                    lease_id=lease_id,
                    conversation_id=conversation_id,
                    data_dir=base_data_dir,
                )
                check_quarantine_path_security(temp_deriv, raw_attachments_root)
                check_quarantine_path_security(deriv_path, raw_attachments_root)

                promote_res = _atomic_no_clobber_promote(temp_deriv, deriv_path, deriv_sha)
                was_newly_created = (promote_res == "fetched")
            except (QuarantineCollisionError, WorkspaceLockError, HashDriftError, SymlinkEscapeError):
                temp_deriv.unlink(missing_ok=True)
                raise
            except Exception as e:
                temp_deriv.unlink(missing_ok=True)
                return {
                    "status": "attachment_conversion_unavailable",
                    "source_sha256": actual_sha,
                    "derivative_sha256": None,
                    "derivative_relative_path": None,
                    "method": "ocr_local_derivative",
                    "tool": "ocrmypdf",
                    "tool_version": _get_tool_version("ocrmypdf"),
                    "quality": "low",
                    "scope": scope,
                    "truncation_reason": truncation_reason,
                    "character_count": 0,
                    "text": "",
                    "error": f"OCR tool execution unavailable: {e}",
                }

            post_source_sha = hashlib.sha256(target_file.read_bytes()).hexdigest().lower()
            if pre_source_sha != post_source_sha:
                if was_newly_created and deriv_path.exists():
                    deriv_path.unlink(missing_ok=True)
                raise RuntimeError("Original quarantine file was mutated during OCR derivative generation!")

            derivative_sha256 = deriv_sha
            derivative_relative_path = f"data/mail-desk/attachments/{run_id}/derivatives/{deriv_filename}"
            method = "ocr_local_derivative"
            tool = "ocrmypdf"
            tool_version = _get_tool_version("ocrmypdf")
            quality = "medium"
            scope["ocr_pages"] = ocr_count

            if pages_to_process > max_ocr_pages:
                truncation_reason = "ocr_page_limit_exceeded"
                ocr_text += f"\n[OCR page limit of {max_ocr_pages} pages exceeded: skipped {pages_to_process - max_ocr_pages} page(s)]"

            if len(ocr_text) > max_chars:
                extracted_text = ocr_text[:max_chars]
                truncation_reason = "max_chars_exceeded"
            else:
                extracted_text = ocr_text

        # Case 2c: Mixed PDF (Digital + Image Pages)
        else:
            # 1. Lock & Path check immediately before temp creation
            active_lock_verifier = lock_verifier or verify_workspace_lock
            active_lock_verifier(
                workspace_root=workspace_root,
                lease_id=lease_id,
                conversation_id=conversation_id,
                data_dir=base_data_dir,
            )

            raw_derivatives_dir = run_dir / "derivatives"
            check_quarantine_path_security(raw_derivatives_dir, raw_attachments_root)
            derivatives_dir = raw_derivatives_dir.resolve()
            try:
                derivatives_dir.relative_to(run_dir)
            except ValueError:
                raise ValueError(f"Derivatives dir '{derivatives_dir}' escapes run directory '{run_dir}'")
            derivatives_dir.mkdir(parents=True, exist_ok=True)

            deriv_filename = f"{target_file.stem}.ocr.pdf"
            raw_deriv_path = derivatives_dir / deriv_filename
            check_quarantine_path_security(raw_deriv_path, raw_attachments_root)
            deriv_path = raw_deriv_path.resolve()
            try:
                deriv_path.relative_to(derivatives_dir)
            except ValueError:
                raise ValueError(f"Derivative path '{deriv_path}' escapes derivatives dir '{derivatives_dir}'")

            img_page_nums = [p[0] for p in image_pages]
            img_pages_to_ocr = img_page_nums[:max_ocr_pages]
            img_pages_over_budget = img_page_nums[max_ocr_pages:]

            pre_source_sha = hashlib.sha256(target_file.read_bytes()).hexdigest().lower()
            if pre_source_sha != actual_sha:
                raise HashDriftError("Source file hash changed prior to OCR execution.")

            temp_deriv = (
                temp_deriv_path
                if temp_deriv_path is not None
                else derivatives_dir / f".{deriv_filename}.{uuid.uuid4().hex}.tmp"
            )
            check_quarantine_path_security(temp_deriv, raw_attachments_root)
            ocr_page_texts: dict[int, str] = {}
            ocr_succeeded = False
            was_newly_created = False

            active_ocr_runner = ocr_runner or _run_ocr_derivative
            try:
                deriv_bytes, _, ocr_count = run_with_timeout(
                    active_ocr_runner,
                    args=(target_file, temp_deriv, max_ocr_pages, img_pages_to_ocr, ocr_timeout),
                    timeout_seconds=ocr_timeout,
                    allow_thread_fallback=False,
                )
                if not temp_deriv.exists():
                    temp_deriv.write_bytes(deriv_bytes)
                deriv_sha = hashlib.sha256(temp_deriv.read_bytes()).hexdigest().lower()

                # 2. Lock & Path check immediately before promotion
                active_lock_verifier(
                    workspace_root=workspace_root,
                    lease_id=lease_id,
                    conversation_id=conversation_id,
                    data_dir=base_data_dir,
                )
                check_quarantine_path_security(temp_deriv, raw_attachments_root)
                check_quarantine_path_security(deriv_path, raw_attachments_root)

                promote_res = _atomic_no_clobber_promote(temp_deriv, deriv_path, deriv_sha)
                was_newly_created = (promote_res == "fetched")
                derivative_sha256 = deriv_sha
                derivative_relative_path = f"data/mail-desk/attachments/{run_id}/derivatives/{deriv_filename}"
                ocr_succeeded = True

                deriv_doc = pymupdf.open(deriv_path)
                for p_num in img_pages_to_ocr:
                    page_idx = p_num - 1
                    if 0 <= page_idx < len(deriv_doc):
                        ocr_page_texts[p_num] = deriv_doc[page_idx].get_text()
                    elif len(deriv_doc) > 0 and len(img_pages_to_ocr) == len(deriv_doc):
                        idx = img_pages_to_ocr.index(p_num)
                        ocr_page_texts[p_num] = deriv_doc[idx].get_text()
                    else:
                        ocr_page_texts[p_num] = ""
                deriv_doc.close()
            except (QuarantineCollisionError, WorkspaceLockError, HashDriftError, SymlinkEscapeError):
                temp_deriv.unlink(missing_ok=True)
                raise
            except Exception:
                temp_deriv.unlink(missing_ok=True)
                ocr_succeeded = False

            post_source_sha = hashlib.sha256(target_file.read_bytes()).hexdigest().lower()
            if pre_source_sha != post_source_sha:
                if was_newly_created and deriv_path.exists():
                    deriv_path.unlink(missing_ok=True)
                raise RuntimeError("Original quarantine file was mutated during OCR derivative generation!")

            lines = []
            cur_chars = 0
            for p_num, has_txt, native_txt, _ in pages_info:
                if has_txt:
                    chunk = native_txt
                elif p_num in img_pages_to_ocr and ocr_succeeded:
                    chunk = ocr_page_texts.get(p_num, "") or f"[OCR Page {p_num}]"
                elif p_num in img_pages_to_ocr and not ocr_succeeded:
                    chunk = f"[Page {p_num}: image page - OCR unavailable]"
                else:
                    chunk = f"[Page {p_num}: image page skipped - OCR page limit of {max_ocr_pages} exceeded]"

                if cur_chars + len(chunk) > max_chars:
                    rem = max_chars - cur_chars
                    lines.append(chunk[:rem])
                    cur_chars += rem
                    truncation_reason = "max_chars_exceeded"
                    break
                lines.append(chunk)
                cur_chars += len(chunk)

            extracted_text = "\n".join(lines)
            method = "mixed_native_ocr" if ocr_succeeded else "native"
            tool = "pymupdf+ocrmypdf" if ocr_succeeded else "pymupdf"
            tool_version = (
                f"pymupdf={_get_tool_version('pymupdf')};ocrmypdf={_get_tool_version('ocrmypdf')}"
                if ocr_succeeded else _get_tool_version("pymupdf")
            )
            quality = "mixed" if ocr_succeeded else "partial"
            scope["ocr_pages"] = len(img_pages_to_ocr) if ocr_succeeded else 0

            if img_pages_over_budget and truncation_reason is None:
                truncation_reason = "ocr_page_limit_exceeded"
            elif not ocr_succeeded and truncation_reason is None:
                truncation_reason = "ocr_unavailable"

    # 3. Microsoft Word (DOCX only - DOCM blocked before)
    elif ext == ".docx":
        import docx
        try:
            doc = docx.Document(target_file)
        except Exception as e:
            return {
                "status": "corrupt_attachment",
                "source_sha256": actual_sha,
                "derivative_sha256": None,
                "derivative_relative_path": None,
                "method": "native",
                "tool": "python-docx",
                "tool_version": _get_tool_version("python-docx"),
                "quality": "low",
                "scope": scope,
                "truncation_reason": None,
                "character_count": 0,
                "text": "",
                "error": f"DOCX parse error: {e}",
            }

        total_paras = len(doc.paragraphs)
        max_paras = limits["max_docx_paragraphs"]
        scope["paragraphs"] = min(total_paras, max_paras)
        if total_paras > max_paras:
            truncation_reason = "max_paragraphs_exceeded"

        lines = []
        cur_chars = 0
        for idx, p in enumerate(doc.paragraphs):
            if idx >= max_paras:
                break
            txt = p.text
            if cur_chars + len(txt) > max_chars:
                rem = max_chars - cur_chars
                lines.append(txt[:rem])
                cur_chars += rem
                truncation_reason = "max_chars_exceeded"
                break
            lines.append(txt)
            cur_chars += len(txt)

        extracted_text = "\n".join(lines)
        method = "native"
        tool = "python-docx"
        tool_version = _get_tool_version("python-docx")

    # 4. Microsoft Excel (XLSX only - XLSM blocked before)
    elif ext == ".xlsx":
        import openpyxl
        try:
            wb = openpyxl.load_workbook(target_file, read_only=True, data_only=True)
        except Exception as e:
            return {
                "status": "corrupt_attachment",
                "source_sha256": actual_sha,
                "derivative_sha256": None,
                "derivative_relative_path": None,
                "method": "native",
                "tool": "openpyxl",
                "tool_version": _get_tool_version("openpyxl"),
                "quality": "low",
                "scope": scope,
                "truncation_reason": None,
                "character_count": 0,
                "text": "",
                "error": f"XLSX parse error: {e}",
            }

        max_sheets = limits["max_xlsx_sheets"]
        max_rows = limits["max_xlsx_rows"]
        max_cols = limits["max_xlsx_cols"]

        total_sheets = len(wb.sheetnames)
        sheets_to_process = min(total_sheets, max_sheets)
        scope["sheets"] = sheets_to_process

        lines = []
        cur_chars = 0
        grid_exceeded = total_sheets > max_sheets

        for sheetname in wb.sheetnames[:sheets_to_process]:
            ws = wb[sheetname]
            sheet_header = f"--- Sheet: {sheetname} ---"
            if cur_chars + len(sheet_header) > max_chars:
                truncation_reason = "max_chars_exceeded"
                break
            lines.append(sheet_header)
            cur_chars += len(sheet_header)

            row_idx = 0
            for row in ws.iter_rows(values_only=True):
                row_idx += 1
                if row_idx > max_rows:
                    grid_exceeded = True
                    break
                if len(row) > max_cols:
                    grid_exceeded = True
                vals = [str(c) if c is not None else "" for c in row[:max_cols]]
                row_str = "\t".join(vals)
                if cur_chars + len(row_str) > max_chars:
                    rem = max_chars - cur_chars
                    lines.append(row_str[:rem])
                    cur_chars += rem
                    truncation_reason = "max_chars_exceeded"
                    break
                lines.append(row_str)
                cur_chars += len(row_str)
            if cur_chars >= max_chars:
                break

        wb.close()
        if truncation_reason is None and grid_exceeded:
            truncation_reason = "grid_limit_exceeded"
        extracted_text = "\n".join(lines)
        method = "native"
        tool = "openpyxl"
        tool_version = _get_tool_version("openpyxl")

    # 5. Microsoft PowerPoint (PPTX only - PPTM blocked before)
    elif ext == ".pptx":
        import pptx
        try:
            prs = pptx.Presentation(target_file)
        except Exception as e:
            return {
                "status": "corrupt_attachment",
                "source_sha256": actual_sha,
                "derivative_sha256": None,
                "derivative_relative_path": None,
                "method": "native",
                "tool": "python-pptx",
                "tool_version": _get_tool_version("python-pptx"),
                "quality": "low",
                "scope": scope,
                "truncation_reason": None,
                "character_count": 0,
                "text": "",
                "error": f"PPTX parse error: {e}",
            }

        max_slides = limits["max_pptx_slides"]
        total_slides = len(prs.slides)
        slides_to_process = min(total_slides, max_slides)
        scope["slides"] = slides_to_process
        if total_slides > max_slides:
            truncation_reason = "max_slides_exceeded"

        lines = []
        cur_chars = 0
        for s_idx in range(slides_to_process):
            slide = prs.slides[s_idx]
            slide_header = f"--- Slide {s_idx + 1} ---"
            if cur_chars + len(slide_header) > max_chars:
                truncation_reason = "max_chars_exceeded"
                break
            lines.append(slide_header)
            cur_chars += len(slide_header)

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text
                        if cur_chars + len(txt) > max_chars:
                            rem = max_chars - cur_chars
                            lines.append(txt[:rem])
                            cur_chars += rem
                            truncation_reason = "max_chars_exceeded"
                            break
                        lines.append(txt)
                        cur_chars += len(txt)
                    if cur_chars >= max_chars:
                        break
            if cur_chars >= max_chars:
                break

        extracted_text = "\n".join(lines)
        method = "native"
        tool = "python-pptx"
        tool_version = _get_tool_version("python-pptx")

    # 6. Fallback: MarkItDown for other supported formats
    else:
        try:
            from markitdown import MarkItDown
            md = MarkItDown()
            res = md.convert(str(target_file))
            raw_text = res.text_content or ""
            if len(raw_text) > max_chars:
                extracted_text = raw_text[:max_chars]
                truncation_reason = "max_chars_exceeded"
            else:
                extracted_text = raw_text
            method = "markitdown"
            tool = "markitdown"
            tool_version = _get_tool_version("markitdown")
        except Exception:
            return {
                "status": "attachment_conversion_unavailable",
                "source_sha256": actual_sha,
                "derivative_sha256": None,
                "derivative_relative_path": None,
                "method": "native",
                "tool": "none",
                "tool_version": "unknown",
                "quality": "low",
                "scope": scope,
                "truncation_reason": None,
                "character_count": 0,
                "text": "",
                "error": f"No extraction converter available for format '{ext}' ({eff_mime})",
            }

    if source_character_count is None and status == "extracted" and truncation_reason is None:
        source_character_count = len(extracted_text)

    return {
        "status": status,
        "source_sha256": actual_sha,
        "derivative_sha256": derivative_sha256,
        "derivative_relative_path": derivative_relative_path,
        "method": method,
        "tool": tool,
        "tool_version": tool_version,
        "quality": quality,
        "scope": scope,
        "truncation_reason": truncation_reason,
        "character_count": len(extracted_text),
        "source_character_count": source_character_count,
        "text": extracted_text,
        "error": None,
    }


def extract_attachment_content(
    fetch_result: dict[str, Any],
    expected_sha256: str,
    data_dir: Path | None = None,
    policy: dict[str, Any] | None = None,
    workspace_root: str | Path | None = None,
    lease_id: str | None = None,
    conversation_id: str | None = None,
    _ocr_runner: Callable[..., Any] | None = None,
    _lock_verifier: Callable[..., Any] | None = None,
) -> dict[str, Any]:
    """Extract bounded textual content from a verified MD-A2 quarantine attachment.

    Enforces strict pre-verification of the MD-A2 envelope, path containment,
    active-format blocking, policy limits, execution timeouts, and derivative security.
    """
    # 1. Strict MD-A2 Envelope Verification (expected_sha256 is mandatory)
    run_id, rel_path, verified_sha, eff_mime = validate_mda2_fetch_result(
        fetch_result, expected_sha256=expected_sha256
    )

    # 2. Resolve Extraction Limits from Policy
    limits = _resolve_extraction_limits(policy)

    # 3. Resolve & Verify Quarantine Path Containment (Preserving unresolved root for symlink checks)
    base_data_dir = data_dir or resolve_data_dir()
    target_file, run_dir, raw_attachments_root = resolve_quarantine_source_file(
        run_id, rel_path, base_data_dir
    )

    if not target_file.exists() or not target_file.is_file():
        raise FileNotFoundError(f"Quarantine file not found on disk: {target_file}")

    file_bytes = target_file.read_bytes()
    actual_sha = hashlib.sha256(file_bytes).hexdigest().lower()

    # 4. Mandatory Pre-Extraction Hash Match against MD-A2 verified hash
    if actual_sha != verified_sha:
        raise HashDriftError(
            f"Quarantine file hash drift: disk file has '{actual_sha}', MD-A2 fetch verified '{verified_sha}'"
        )

    # 5. Fail-Closed Active Content & Disallowed Extension Enforcement
    ext = target_file.suffix.lower()
    disallowed_exts = limits["disallowed_extensions"]

    if ext in disallowed_exts:
        raise DisallowedExtensionError(f"Disallowed active extension blocked: '{ext}'")

    # Sniff magic bytes & inspect package for active code/macros
    sniffed_mime, is_active, active_reason = detect_mime_and_active_content(
        target_file.name, file_bytes, policy=policy or DEFAULT_ATTACHMENT_POLICY
    )
    if is_active or sniffed_mime == "application/vnd.ms-office.active-macro":
        raise ActiveContentBlockedError(
            f"Active content blocked in attachment '{target_file.name}': {active_reason or sniffed_mime}"
        )

    # 6. Re-validate Effective MIME and Extension Contract against MD-A2 Stated MIME
    validate_mime_and_extension(
        effective_mime=sniffed_mime,
        inventory_mime=eff_mime,
        filename=target_file.name,
        policy=policy or DEFAULT_ATTACHMENT_POLICY,
    )
    # Check exact match if not octet-stream
    if eff_mime != "application/octet-stream" and sniffed_mime != eff_mime:
        raise MimeDriftError(
            f"Effective MIME drift detected: MD-A2 declared '{eff_mime}', but content sniffed as '{sniffed_mime}'"
        )

    # 7. Execute bounded extraction under interruptible process timeout
    process_timeout = limits["process_timeout_seconds"]

    # Pre-determine invocation-specific temp derivative path for parent-owned zero-leakage cleanup
    raw_derivatives_dir = run_dir / "derivatives"
    check_quarantine_path_security(raw_derivatives_dir, raw_attachments_root)
    derivatives_dir = raw_derivatives_dir.resolve()
    try:
        derivatives_dir.relative_to(run_dir)
    except ValueError:
        raise ValueError(f"Derivatives dir '{derivatives_dir}' escapes run directory '{run_dir}'")

    invocation_id = uuid.uuid4().hex
    deriv_filename = f"{target_file.stem}.ocr.pdf"
    known_temp_deriv_path = derivatives_dir / f".{deriv_filename}.{invocation_id}.tmp"

    try:
        res = run_with_timeout(
            _extract_content_internal,
            args=(
                target_file,
                actual_sha,
                target_file.name,
                ext,
                eff_mime,
                run_dir,
                run_id,
                raw_attachments_root,
                base_data_dir,
                limits,
                workspace_root,
                lease_id,
                conversation_id,
                known_temp_deriv_path,
                _ocr_runner,
                _lock_verifier,
            ),
            timeout_seconds=process_timeout,
            allow_thread_fallback=False,
        )
        if isinstance(res, dict) and res.get("status") != "extracted":
            _cleanup_temp_artifacts(known_temp_deriv_path, derivatives_dir, target_file.stem)
        return res
    except TimeoutError as err:
        _cleanup_temp_artifacts(known_temp_deriv_path, derivatives_dir, target_file.stem)
        return {
            "status": "extraction_failed",
            "source_sha256": actual_sha,
            "derivative_sha256": None,
            "derivative_relative_path": None,
            "method": "unknown",
            "tool": "unknown",
            "tool_version": "unknown",
            "quality": "low",
            "scope": {
                "pages_processed": None,
                "pages_total": None,
                "ocr_pages": 0,
                "paragraphs": None,
                "slides": None,
                "sheets": None,
            },
            "truncation_reason": "timeout_exceeded",
            "character_count": 0,
            "text": "",
            "error": f"Attachment extraction timed out after {process_timeout}s: {err}",
        }
    except Exception:
        _cleanup_temp_artifacts(known_temp_deriv_path, derivatives_dir, target_file.stem)
        raise
